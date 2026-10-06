"""Build the Age of Origins upgrade tables (CSV + XLSX) from the source snapshots in ../sources.

Buildings are reconciled from two independent sources:
  * aoztips.com building calculator (current data, includes uranium and levels 39-41, no build times)
  * appmedia.jp Age of Z wiki (2021-22 snapshot, levels up to 38, includes base build times,
    prerequisites and per-level effects; also covers resource/support buildings aoztips lacks)
Every building row carries a Source and a Flag column saying how the numbers were arrived at.
Everything else (officers, equipment, titans, warplanes, antiques, merit) comes from aoztips.com.

Run: python3 age_of_origins/scripts/build_tables.py
"""
import csv
import json
import re
from collections import defaultdict
from pathlib import Path

import progression
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "sources"
OUT = ROOT / "output"
RES = ["food", "oil", "steel", "minerals"]

# appmedia (JP) building name -> English in-game name. Names marked * are not confirmed in English sources.
JP_NAMES = {
    "メインホール": "Main Hall", "防御壁": "City Wall", "防衛壁": "City Wall", "防御塔": "Defense Tower",
    "倉庫": "Depot", "情報センター": "Recon Center", "外交センター": "Embassy", "指令センター": "Dispatch Center",
    "病院": "Hospital", "リハビリ拠点": "Rehab Center*", "兵舎": "Camp", "テクノロジーセンター": "Academy",
    "テクノロジー作業場": "Workshop", "農場": "Farm", "石油精製工場": "Oil Refinery", "鉄鉱山": "Steel Mill",
    "レアアース鉱山": "Mineral Mine", "訓練場": "Training Grounds", "車両工場": "Garage", "機械工場": "Factory",
    "住居": "House", "後方支援センター": "Hospital",
}
# aoztips splits the Recon Center at level 31; everything else uses the in-game name already.
AOZ_NAMES = {"Recon": "Recon Center"}

EFFECT_LABELS = {
    "採集速度": "Gathering speed", "負傷兵最大ユニット数": "Wounded capacity (units)", "ストレージ量": "Storage",
    "生産量": "Production / hour", "最大住民数": "Max residents", "レアアース保護": "Minerals protected",
    "石油保護": "Oil protected", "鉄鉱保護": "Steel protected", "食糧保護": "Food protected",
    "同盟メンバー最大駐屯数": "Max alliance reinforcement (units)", "救援回数": "Alliance helps",
    "兵士ユニット集結上限": "Rally capacity (units)", "治療速度": "Healing speed", "兵士募集加速": "Training speed",
    "最大募集数": "Max units per training batch", "最大部隊兵士ユニット": "Max units per march (fleet size)",
    "ダメージ出力": "Damage", "攻撃範囲": "Attack range", "攻撃間隔": "Attack interval", "都市防衛力": "City defense",
}
RECON_TEXT = {
    "あなたに向かってくる部隊の指揮官名を表示します。": "Shows the commander name of incoming marches",
    "あなたに向かってくる部隊の目的地を表示します。": "Shows the destination of incoming marches",
    "相手城内の具体的な資源量と未収穫量を偵察することができます。": "Scout: exact resources and uncollected resources in target city",
    "だれがその部隊を所有しているか世界地図上で確認することができます。": "Shows on the world map who owns a march",
    "あなたに向かってくる部隊の到着時間を表示します。": "Shows arrival time of incoming marches",
    "あなたに向かってくる部隊がどの座標から来ているかを表示します。": "Shows origin coordinates of incoming marches",
    "あなたに向かってくる部隊のおよそのユニット数を表示します。": "Shows approximate unit count of incoming marches",
    "相手の守備部隊のおよそのユニット数を偵察できます。": "Scout: approximate garrison unit count",
    "あなたに向かってくる部隊の指揮官の所属都市レベルを表示します。": "Shows city level of incoming commander",
    "相手防御塔レベルを偵察できます。": "Scout: defense tower levels",
    "あなたに向かってくる兵種とおよその数を教えてく表示します。": "Shows troop types and approximate numbers of incoming marches",
    "相手の守備部隊の兵種とおよその数を偵察できます。": "Scout: garrison troop types and approximate numbers",
    "相手支援部隊のおよそのユニット数を偵察できます。": "Scout: approximate reinforcement unit count",
    "あなたに向かってくる全兵種と具体的な規模を表示します。": "Shows all troop types and exact sizes of incoming marches",
    "相手支援部隊の兵種と都市レベルを偵察できます。": "Scout: reinforcement troop types and city level",
    "相手の防御壁の耐久性最大値と現在地を偵察できます。": "Scout: city wall max and current durability",
    "相手支援部隊の兵種とおよその規模を偵察できます。": "Scout: reinforcement troop types and approximate size",
    "敵方のタイタン情報を偵察できます": "Scout: enemy Titan info",
    "相手守衛部隊の兵種と具体的な規模を偵察できます。": "Scout: garrison troop types and exact size",
    "偵察行軍速度＋10％": "Scout march speed +10%",
    "相手支援部隊の全兵種と具体的な規模を偵察できます。": "Scout: all reinforcement troop types and exact size",
    "あなたに向かてくる部隊のテクノロジーセンターレベルを表示します。": "Shows Academy level of incoming commander",
    "相手の装置情報を偵察することができます。": "Scout: target's device (gear) info",
    "相手のテクノロジーセンターレベルを偵察できます。": "Scout: target's Academy level",
    "あなたに向かってくる部隊の指揮官スキルを表示します。": "Shows commander skills of incoming marches",
    "相手の副官情報を偵察することができます。": "Scout: target's officers",
    "相手のキャラテクノロジーを偵察できます": "Scout: target's commander tech",
    "ターゲットの生物化学技術レベルを偵察できます。": "Scout: target's biochemical tech level",
    "偵察速度15％アップ": "Scout speed +15%",
    "相手の副官スキル情報を偵察できます。": "Scout: target's officer skills",
    "偵察速度25％アップ": "Scout speed +25%",
    "ターゲットの拡張機器効果を偵察できます。": "Scout: target's expansion-gear effects",
}
# Buildings that legitimately share a cost table (troop-training buildings).
SAME_COST = {"Camp": ("Factory",), "Factory": ("Camp",)}
TIER_WORDS = {"初級": "Basic ", "精鋭": "Elite ", "中級": "Intermediate "}
OTHER_ITEMS = {"砲塔パーツ": "Turret Parts"}


def load(path):
    return json.loads((SRC / path).read_text(encoding="utf-8"))


def fmt_time(sec):
    if sec is None:
        return ""
    d, rem = divmod(sec, 86400)
    h, rem = divmod(rem, 3600)
    m, s = divmod(rem, 60)
    return f"{d}d {h:02}:{m:02}:{s:02}" if d else f"{h:02}:{m:02}:{s:02}"


def translate_effect(text):
    m = re.match(r"(.+?)：(.+)", text)
    if m:
        return f"{EFFECT_LABELS.get(m.group(1), m.group(1))}: {m.group(2)}"
    if text in RECON_TEXT:
        return RECON_TEXT[text]
    if text == "-":
        return ""
    for jp, en in TIER_WORDS.items():
        if text.startswith(jp):
            return f"Unlocks {en.strip()} unit ({text})"
    return f"Unlocks unit ({text})"  # higher-tier troop names: official English names not verified


def translate_req(req):
    m = re.match(r"(.+?)Lv(\d+)", req)
    return f"{JP_NAMES.get(m.group(1), m.group(1))} {m.group(2)}" if m else req


def core(res):
    return {k: res.get(k, 0) or 0 for k in RES}


def pick(field, level, a_val, j_val, agreed, base):
    """Resolve a genuine conflict: keep the candidate whose ratio to the level's food/oil best matches neighbours."""
    def ratio_at(lv):
        r = agreed.get(lv)
        base = r and (r["food"] or r["oil"])
        return r[field] / base if base and r[field] else None

    def nearest(step):
        for d in range(1, 6):
            x = ratio_at(level + step * d)
            if x:
                return x
    neigh = [x for x in (nearest(-1), nearest(1)) if x]
    if neigh and base and field not in ("food", "oil"):
        target = sum(neigh) / len(neigh)
        return min((a_val, "aoztips"), (j_val, "appmedia"), key=lambda c: abs(c[0] / base - target))
    # Single-resource line: prefer the value closer to the geometric mean of its neighbours.
    nb = [agreed[lv][field] for lv in (level - 1, level + 1) if lv in agreed and agreed[lv][field]]
    if len(nb) == 2:
        gm = (nb[0] * nb[1]) ** 0.5
        return min((a_val, "aoztips"), (j_val, "appmedia"), key=lambda c: abs(c[0] - gm))
    return (a_val, "aoztips")


def build_buildings():
    aoz = defaultdict(dict)
    for r in load("aoztips/buildings.json"):
        aoz[AOZ_NAMES.get(r["name"], r["name"])][r["level"]] = r
    jp = {}
    jp_all = {k: v for k, v in load("appmedia/facilities.json").items() if v["costs"]}
    for jname, rec in jp_all.items():
        if rec["costs"]:
            jp[JP_NAMES[jname]] = (jname, rec)

    rows, conflicts = [], []
    for name in sorted(set(aoz) | set(jp)):
        a = aoz.get(name, {})
        jname, jrec = jp.get(name, (None, None))
        jc = {int(k): v for k, v in (jrec or {}).get("costs", {}).items()}
        je = {int(k): v for k, v in (jrec or {}).get("effects", {}).items()}

        def a_core(lv):
            r = a.get(lv)
            return core(r["resources"]) if r else None

        # Detect rows aoztips lists one level too high (its level L+1 equals appmedia's level L).
        shifted = {lv for lv in jc if a_core(lv + 1) and any(a_core(lv + 1).values())
                   and a_core(lv + 1) == core(jc[lv]["resources"]) and a_core(lv) != core(jc[lv]["resources"])}
        comparable = [lv for lv in jc if a_core(lv) is not None]
        whole_shift = bool(comparable) and len(shifted) >= 0.8 * len(comparable)

        # appmedia rows that are verbatim copies of another building's row at the same level (copy-paste errors).
        copied = {lv for lv in jc if any(core(jc[lv]["resources"]).values()) and any(
            o != jname and str(lv) in r["costs"] and core(r["costs"][str(lv)]["resources"]) == core(jc[lv]["resources"])
            and JP_NAMES[o] not in SAME_COST.get(name, ()) for o, r in jp_all.items())}
        # Levels agreed by both sources feed the ratio checks used to settle conflicts.
        agreed = {lv: core(jc[lv]["resources"]) for lv in jc
                  if a_core(lv + 1 if (whole_shift or lv in shifted) else lv) == core(jc[lv]["resources"])}
        levels = sorted(set(jc) | ({lv - 1 for lv in a} if whole_shift else set(a)))
        for lv in levels:
            if lv < 1:
                continue
            a_lv = lv + 1 if (whole_shift or lv in shifted) else lv
            ar = a.get(a_lv)
            j = jc.get(lv)
            a_res = core(ar["resources"]) if ar else None
            a_blank = a_res is not None and not any(a_res.values()) and j and any(core(j["resources"]).values())
            flag, source = [], []
            if j and a_res and not a_blank:
                jres = core(j["resources"])
                if a_lv != lv:
                    flag.append(f"aoztips row offset corrected (its level {a_lv} = this level)")
                final = dict(jres)
                prev = jc.get(lv - 1)
                if a_lv == lv and a_res != jres and prev and a_res == core(prev["resources"]):
                    # aoztips repeats the previous level's cost here (tail end of an offset run).
                    flag.append(f"aoztips row duplicates level {lv - 1}; appmedia used")
                    a_res = jres
                for f in RES:
                    if a_res[f] != jres[f]:
                        if lv in copied:
                            val, who = a_res[f], "aoztips"
                        else:
                            base = (a_res["food"] or a_res["oil"]) if a_res["food"] == jres["food"] and a_res["oil"] == jres["oil"] else None
                            val, who = pick(f, lv, a_res[f], jres[f], agreed, base)
                        final[f] = val
                        conflicts.append([name, lv, f, a_res[f], jres[f], val, who])
                        flag.append(f"{f} conflict: aoztips {a_res[f]:,} vs appmedia {jres[f]:,}; kept {who}"
                                    + (" (appmedia row copies another building)" if lv in copied else ""))
                source.append("aoztips+appmedia")
                if not flag:
                    flag.append("verified: 2 sources agree")
            elif j:
                final = core(j["resources"])
                source.append("appmedia")
                flag.append("single source (appmedia 2021-22)" + ("; aoztips row blank" if a_blank else ""))
            elif ar:
                final = a_res
                source.append("aoztips")
                flag.append("single source (aoztips 2026)")
            else:
                continue
            uranium = (ar or {}).get("resources", {}).get("uranium", 0) or (j or {}).get("resources", {}).get("uranium", 0) or 0
            badges = (ar or {}).get("resources", {}).get("leader_badges", 0) or (j or {}).get("resources", {}).get("leader_badges", 0) or 0
            other = []
            for o in (j or {}).get("resources", {}).get("_other", []):
                k, _, v = o.partition(" / ")
                other.append(f"{OTHER_ITEMS.get(k, k)} x{v}")
            if j:
                req = "; ".join(translate_req(x) for x in j["requires"])
            elif ar:
                req = "; ".join(x for x in ar["requirements"].values() if x)
            else:
                req = ""
            effects = "; ".join(e for e in (translate_effect(x) for x in je.get(lv, [])) if e)
            if name.endswith("*"):
                flag.append("English building name not confirmed")
            if name == "Core Mining Station":
                flag.append("unverified: single odd row in aoztips")
            rows.append({
                "Building": name.rstrip("*"), "JP name": jname or "", "Level": lv,
                "Food": final["food"], "Oil": final["oil"], "Steel": final["steel"], "Minerals": final["minerals"],
                "Uranium": uranium, "Leader Badges": badges, "Other items": ", ".join(other),
                "Base build time": fmt_time(j["time_seconds"]) if j else "",
                "Build time (hours)": round(j["time_seconds"] / 3600, 2) if j and j["time_seconds"] is not None else "",
                "Prerequisites": req, "Effect at this level": effects,
                "Source": "+".join(source), "Flag": " | ".join(flag),
            })
    return rows, conflicts


def building_totals(rows):
    out = {}
    for r in rows:
        t = out.setdefault(r["Building"], {"Building": r["Building"], "Levels covered": "", "Max level": 0,
                                           "Food": 0, "Oil": 0, "Steel": 0, "Minerals": 0, "Uranium": 0,
                                           "Leader Badges": 0, "Base build time, levels <=38 (days)": 0.0, "_min": 999})
        for k in ("Food", "Oil", "Steel", "Minerals", "Uranium", "Leader Badges"):
            t[k] += r[k]
        if r["Build time (hours)"] != "":
            t["Base build time, levels <=38 (days)"] += r["Build time (hours)"] / 24
        t["Max level"] = max(t["Max level"], r["Level"])
        t["_min"] = min(t["_min"], r["Level"])
    for t in out.values():
        t["Levels covered"] = f'{t.pop("_min")}-{t["Max level"]}'
        t["Base build time, levels <=38 (days)"] = round(t["Base build time, levels <=38 (days)"], 1)
    return list(out.values())


def merit_ranks():
    return [{"Merit rank": r["rank"], "Attack %": r["attack"], "HP %": r["hp"], "Fleet (units)": r["fleet"]}
            for r in load("aoztips/merit_ranks.json")]


GAPS = [
    ("Academy research (all branches: Development, Military, Resource, Defense, etc.)",
     "No accessible public source publishes per-level research costs/values. aozcodes.com has research pages but is behind a Cloudflare bot check this environment cannot pass."),
    ("Troop tiers (training cost, stats, upkeep per T1-T12 unit)", "Not found in an accessible structured source."),
    ("Resource buildings (Farm, Oil Refinery, Steel Mill, Mineral Mine), Training Grounds, House, Rehab Center, Defense Tower",
     "Only the 2021-22 appmedia.jp snapshot (max level 30 / 10). Levels added since then are missing; costs may have been rebalanced."),
    ("Base build times", "From appmedia.jp (2021-22), and only up to level 38. Times are before construction-speed buffs. aoztips publishes no times."),
    ("Uranium costs", "Only from aoztips; appmedia predates uranium. Levels 31+ in appmedia-only rows show 0 uranium, which may be wrong."),
    ("Biochemical Lab, Elite Equipment Facility, Alliance tech, VIP, Commander skills, Battle machines, Mega devices",
     "No per-level numeric source found."),
]


def write(sheets, filename, csv_dir):
    OUT.mkdir(exist_ok=True)
    (OUT / csv_dir).mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="305496")
    for title, rows in sheets:
        if not rows:
            continue
        cols = list(rows[0].keys())
        with open(OUT / csv_dir / f"{re.sub(r'[^a-z0-9]+', '_', title.lower()).strip('_')}.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=cols)
            w.writeheader()
            w.writerows(rows)
        ws = wb.create_sheet(title[:31])
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head, fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append([r[c] for c in cols])
        for i, c in enumerate(cols, 1):
            width = max([len(str(c))] + [len(str(r[c])) for r in rows[:400]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 70)
            if all(isinstance(r[c], int) for r in rows if r[c] != ""):
                for cell in ws[get_column_letter(i)][1:]:
                    cell.number_format = "#,##0"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(OUT / filename)


def main():
    b_rows, conflicts = build_buildings()
    conf_rows = [{"Building": c[0], "Level": c[1], "Resource": c[2], "aoztips value": c[3], "appmedia value": c[4],
                  "Value kept": c[5], "Kept from": c[6]} for c in conflicts]
    gaps = [{"Upgrade system": g, "Status / reason": r} for g, r in GAPS]
    books = [
        ("age_of_origins_buildings.xlsx", "csv/buildings", [
            ("Buildings", b_rows), ("Building totals", building_totals(b_rows)), ("Building conflicts", conf_rows),
            ("Merit ranks", merit_ranks()), ("Gaps - not covered", gaps)]),
        ("age_of_origins_officers_titans_warplanes_antiques.xlsx", "csv/progression", progression.build()),
    ]
    for filename, csv_dir, sheets in books:
        write(sheets, filename, csv_dir)
        print(filename)
        for t, r in sheets:
            print(f"  {t:32} {len(r):6} rows")


if __name__ == "__main__":
    main()
